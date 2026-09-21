"""In-memory plan-constructor contracts and planner adapter.

The constructor deliberately never writes the source Excel or processed CSV files.
Browser drafts arrive as JSON, are normalized into a temporary planning context,
and disappear after the request has been served.
"""

from __future__ import annotations

import hashlib
import tempfile
from copy import deepcopy
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Literal
from zipfile import BadZipFile

from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, Field

from src.data.normalization import NormalizedDataset, normalize
from src.data.parser import DatasetParsingError, parse_workbook
from src.data.validation import validate
from src.skills import (
    EXTRACTOR_VERSION,
    canonicalize_skill,
    get_skill_registry,
    merge_task_skill_requirements,
)

from .service import TABLE_FILES, PlanningAPIService, ServiceDataError
from .catalog import CatalogRepository
from .data_preflight import build_data_preflight
from src.planning.task_states import dependency_is_satisfied, is_historical_done


MAX_CONSTRUCTOR_UPLOAD_BYTES = 25 * 1024 * 1024


class ConstructorTeam(BaseModel):
    team_id: str = Field(min_length=1, max_length=120)
    name: str | None = Field(default=None, max_length=200)
    # Existing teams provide historical velocity. A newly created team has no
    # history yet and may provide a direct per-sprint capacity instead.
    expected_velocity_sp: float | None = Field(default=None, ge=0)
    capacity_sp: float | None = Field(default=None, ge=0)
    focus_factor: float = Field(default=0.8, ge=0, le=1)


class ConstructorEngineerAssignment(BaseModel):
    # ``None`` denotes an unassigned/shared-pool engineer. Such a resource is
    # available to help any selected team during manual planning.
    team_id: str | None = Field(default=None, max_length=120)
    role: str = Field(min_length=1, max_length=200)
    grade: str | None = Field(default=None, max_length=120)
    capacity_rate: float = Field(gt=0, le=1)


class ConstructorEngineer(BaseModel):
    engineer_id: str = Field(min_length=1, max_length=120)
    name: str | None = Field(default=None, max_length=200)
    assignments: list[ConstructorEngineerAssignment] = Field(min_length=1)
    skills: list[str] = Field(default_factory=list)


class ConstructorRoleRequirement(BaseModel):
    role: str = Field(min_length=1, max_length=200)
    required_hours: float = Field(gt=0)
    skills: list[str] = Field(default_factory=list)


class ConstructorSkillRequirement(BaseModel):
    technology: str = Field(min_length=1, max_length=200)
    role: str | None = Field(default=None, max_length=200)
    source: Literal["explicit", "title", "both"] = "explicit"
    confidence: Literal["high", "medium", "low"] = "high"
    matched_text: str | None = Field(default=None, max_length=500)
    confirmed: bool = True
    rejected: bool = False


class ConstructorTask(BaseModel):
    task_id: str = Field(min_length=1, max_length=120)
    title: str = Field(min_length=1, max_length=500)
    initiative_id: str | None = Field(default=None, max_length=120)
    team_id: str | None = Field(default=None, max_length=120)
    status: Literal["ToDo", "InProgress", "Done"] = "ToDo"
    rung: float = 0
    story_points: float = Field(ge=0)
    planned_start: str | None = Field(default=None, max_length=50)
    deadline: str | None = Field(default=None, max_length=50)
    role_requirements: list[ConstructorRoleRequirement] = Field(default_factory=list)
    required_skills: list[ConstructorSkillRequirement] = Field(default_factory=list)
    depends_on: list[str] = Field(default_factory=list)


class ConstructorScenario(BaseModel):
    teams: list[ConstructorTeam] = Field(min_length=1)
    engineers: list[ConstructorEngineer] = Field(min_length=1)
    tasks: list[ConstructorTask] = Field(min_length=1)


class ScenarioValidationError(ValueError):
    """The submitted draft is structurally valid JSON but not a valid scenario."""

    def __init__(self, errors: Iterable[str]):
        self.errors = tuple(errors)
        super().__init__("; ".join(self.errors))


class ConstructorDatasetUploadError(ValueError):
    """An uploaded workbook cannot be used as a constructor catalog."""


@dataclass(frozen=True)
class PreparedConstructorUpload:
    """Parsed and audited upload retained until the user accepts warnings."""

    filename: str
    size_bytes: int
    fingerprint: str
    tables: dict[str, list[dict[str, Any]]]
    validation_warnings: tuple[str, ...]
    audit: dict[str, Any]


ARCHIVE_DECISIONS = {"KEEP_ARCHIVED", "RESTORE", "RESTORE_UPDATE"}
EXISTING_DECISIONS = {"KEEP_CURRENT", "UPDATE"}


def _deduplicate_normalized_dataset(
    dataset: NormalizedDataset,
) -> tuple[NormalizedDataset, dict[str, Any]]:
    """Remove repeated Excel rows deterministically before validation.

    Entity IDs are strict for tasks.  Multi-team/multi-role engineer rows stay
    valid and are deduplicated by their full assignment identity.
    """

    identities = {
        "tasks.csv": ("task_id",),
        "task_role_effort.csv": ("task_id", "role"),
        "task_dependencies.csv": ("blocking_task_id", "blocked_task_id", "dependency_type"),
        "engineers.csv": ("engineer_id", "team_id", "role"),
        "engineer_skills.csv": ("engineer_id", "technology"),
        "task_skill_requirements.csv": ("task_id", "role", "technology"),
        "team_history.csv": ("team_id", "snapshot_date"),
        "planning_results.csv": ("value",),
        "plan_mismatch_reasons.csv": ("value",),
        "task_closure_results.csv": ("value",),
    }
    cleaned: dict[str, list[dict[str, Any]]] = {}
    details: list[dict[str, Any]] = []
    removed_total = 0
    for table_name, rows in dataset.tables.items():
        fields = identities.get(table_name)
        if not fields:
            cleaned[table_name] = [dict(row) for row in rows]
            continue
        latest: dict[tuple[str, ...], tuple[int, dict[str, Any]]] = {}
        unique_without_identity: list[tuple[int, dict[str, Any]]] = []
        duplicate_ids: set[str] = set()
        for index, source in enumerate(rows):
            row = dict(source)
            key = tuple(_clean_text(row.get(field)) for field in fields)
            if not any(key):
                unique_without_identity.append((index, row))
                continue
            if key in latest:
                duplicate_ids.add(" / ".join(value for value in key if value))
                removed_total += 1
            latest[key] = (index, row)
        combined = [*unique_without_identity, *latest.values()]
        cleaned[table_name] = [row for _, row in sorted(combined, key=lambda item: item[0])]
        removed = len(rows) - len(cleaned[table_name])
        if removed:
            details.append({
                "table": table_name,
                "removed": removed,
                "identities": sorted(duplicate_ids),
            })
    return NormalizedDataset(
        tables=cleaned,
        effort_task_ids=set(dataset.effort_task_ids),
        effort_roles=set(dataset.effort_roles),
        placeholder_sources=dict(dataset.placeholder_sources),
    ), {"removed_total": removed_total, "tables": details}


def _validated_archive_decisions(
    prepared: PreparedConstructorUpload,
    decisions: dict[str, str] | None,
) -> dict[str, str]:
    """Normalize explicit archive choices and reject dangling references."""

    records = prepared.audit.get("archived_entities") or []
    known = {str(item.get("decision_key")) for item in records}
    supplied = {str(key): str(value).upper() for key, value in (decisions or {}).items()}
    if set(supplied) - known:
        raise ConstructorDatasetUploadError("Решение относится к неизвестной архивной записи.")
    if any(value not in ARCHIVE_DECISIONS for value in supplied.values()):
        raise ConstructorDatasetUploadError("Выберите допустимое действие для архивной записи.")
    result = {key: supplied.get(key, "KEEP_ARCHIVED") for key in known}
    kept_teams = {
        key.split(":", 1)[1]
        for key, value in result.items()
        if key.startswith("team:") and value == "KEEP_ARCHIVED"
    }
    kept_tasks = {
        key.split(":", 1)[1]
        for key, value in result.items()
        if key.startswith("task:") and value == "KEEP_ARCHIVED"
    }
    for row in prepared.tables.get("tasks", []):
        if _clean_text(row.get("team_id")) in kept_teams and not is_historical_done(row):
            raise ConstructorDatasetUploadError(
                f"Задача {_clean_text(row.get('task_id'))} ссылается на архивную команду "
                f"{_clean_text(row.get('team_id'))}. Восстановите команду или измените команду задачи."
            )
    for row in prepared.tables.get("engineers", []):
        if _clean_text(row.get("team_id")) in kept_teams:
            raise ConstructorDatasetUploadError(
                f"Инженер {_clean_text(row.get('engineer_id'))} ссылается на архивную команду "
                f"{_clean_text(row.get('team_id'))}. Восстановите команду или измените команду инженера."
            )
    active_tasks = {
        _clean_text(row.get("task_id"))
        for row in prepared.tables.get("tasks", [])
        if not is_historical_done(row)
    }
    for row in prepared.tables.get("task_dependencies", []):
        blocking = _clean_text(row.get("blocking_task_id"))
        blocked = _clean_text(row.get("blocked_task_id"))
        if blocking in kept_tasks and blocked in active_tasks:
            raise ConstructorDatasetUploadError(
                f"Задача {blocked} зависит от архивной задачи {blocking}, оставленной в архиве. "
                "Восстановите задачу или исправьте зависимость."
            )
    return result


def _validated_existing_decisions(
    prepared: PreparedConstructorUpload,
    decisions: dict[str, str] | None,
) -> dict[str, str]:
    records = prepared.audit.get("existing_entities") or []
    known = {str(item.get("decision_key")) for item in records}
    supplied = {str(key): str(value).upper() for key, value in (decisions or {}).items()}
    if set(supplied) - known:
        raise ConstructorDatasetUploadError("Решение относится к неизвестной существующей записи.")
    if any(value not in EXISTING_DECISIONS for value in supplied.values()):
        raise ConstructorDatasetUploadError("Выберите допустимое действие для существующей записи.")
    return {key: supplied.get(key, "KEEP_CURRENT") for key in known}


def _clean_text(value: Any) -> str:
    return str(value or "").strip()


def _duplicates(values: Iterable[str]) -> set[str]:
    seen: set[str] = set()
    duplicates: set[str] = set()
    for value in values:
        if value in seen:
            duplicates.add(value)
        seen.add(value)
    return duplicates


def _cycle_nodes(
    tasks: list[ConstructorTask],
    effective_task_rows: dict[str, dict[str, Any]] | None = None,
) -> tuple[str, ...]:
    """Return a deterministic dependency cycle witness, if one exists."""

    rows = effective_task_rows or {}
    active_ids = {
        task.task_id
        for task in tasks
        if not is_historical_done(rows.get(task.task_id, {"status": task.status}))
    }
    graph = {
        task.task_id: tuple(dependency for dependency in task.depends_on if dependency in active_ids)
        for task in tasks
        if task.task_id in active_ids
    }
    visiting: set[str] = set()
    visited: set[str] = set()
    path: list[str] = []

    def visit(task_id: str) -> tuple[str, ...]:
        if task_id in visiting:
            start = path.index(task_id)
            return tuple((*path[start:], task_id))
        if task_id in visited:
            return ()
        visiting.add(task_id)
        path.append(task_id)
        for dependency_id in graph.get(task_id, ()):
            cycle = visit(dependency_id)
            if cycle:
                return cycle
        path.pop()
        visiting.remove(task_id)
        visited.add(task_id)
        return ()

    for task_id in sorted(graph):
        cycle = visit(task_id)
        if cycle:
            return cycle
    return ()


def validate_constructor_scenario(
    scenario: ConstructorScenario,
    catalog_tables: dict[str, list[dict[str, Any]]] | None = None,
) -> None:
    """Validate cross-entity rules that cannot be expressed by field types."""

    errors: list[str] = []
    team_ids = [_clean_text(team.team_id) for team in scenario.teams]
    engineer_ids = [_clean_text(item.engineer_id) for item in scenario.engineers]
    task_ids = [_clean_text(item.task_id) for item in scenario.tasks]

    for entity, values in (
        ("команд", team_ids),
        ("инженеров", engineer_ids),
        ("задач", task_ids),
    ):
        for duplicate in sorted(_duplicates(values)):
            errors.append(f"Дублируется ID {entity}: {duplicate}.")
        if any(not value for value in values):
            errors.append(f"ID {entity} не может быть пустым.")

    known_teams = set(team_ids)
    selected_tasks = set(task_ids)
    catalog_engineer_teams: dict[str, set[str]] = defaultdict(set)
    for row in (catalog_tables or {}).get("engineers", []):
        engineer_id = _clean_text(row.get("engineer_id"))
        team_id = _clean_text(row.get("team_id"))
        if engineer_id and team_id:
            catalog_engineer_teams[engineer_id].add(team_id)
    catalog_task_rows = {
        _clean_text(row.get("task_id")): row
        for row in (catalog_tables or {}).get("tasks", [])
        if _clean_text(row.get("task_id"))
    }
    effective_task_rows = dict(catalog_task_rows)
    for item in scenario.tasks:
        existing = catalog_task_rows.get(item.task_id, {})
        # A historical completion is a fact: neither a stale active catalog
        # row nor a conflicting scenario edit may reactivate it.
        status = "Done" if (
            item.status == "Done" or is_historical_done(existing)
        ) else item.status
        effective_task_rows[item.task_id] = {**existing, "status": status}
    active_scenario_task_ids = {
        task_id
        for task_id in selected_tasks
        if not is_historical_done(effective_task_rows.get(task_id, {}))
    }
    for team in scenario.teams:
        if team.expected_velocity_sp is None and team.capacity_sp is None:
            errors.append(
                f"Команда {team.team_id}: укажите исторический velocity "
                "или ручную capacity за спринт."
            )
        if abs(team.focus_factor - 0.8) > 1e-9:
            errors.append(
                f"Команда {team.team_id}: Focus Factor должен быть 0.8 "
                "по правилам кейса."
            )

    for engineer in scenario.engineers:
        assignment_teams: list[str] = []
        for assignment in engineer.assignments:
            assignment_teams.append(_clean_text(assignment.team_id))
            if not _clean_text(assignment.role):
                errors.append(f"Инженер {engineer.engineer_id}: роль не может быть пустой.")
            if assignment.team_id and assignment.team_id not in known_teams:
                errors.append(
                    f"Инженер {engineer.engineer_id}: команда "
                    f"{assignment.team_id} не добавлена в сценарий."
                )
            source_teams = catalog_engineer_teams.get(engineer.engineer_id, set())
            if source_teams and assignment.team_id and assignment.team_id not in source_teams:
                errors.append(
                    f"Инженер {engineer.engineer_id} уже закреплён за другой "
                    "командой и не может быть перенесён в сценарии."
                )
        for duplicate in sorted(_duplicates(assignment_teams)):
            errors.append(
                f"Инженер {engineer.engineer_id}: команда {duplicate} "
                "указана более одного раза."
            )
        total_capacity_rate = sum(item.capacity_rate for item in engineer.assignments)
        if total_capacity_rate > 1 + 1e-9:
            errors.append(
                f"Инженер {engineer.engineer_id}: сумма Capacity Rate по командам "
                f"равна {total_capacity_rate:g}, но не может превышать 1."
            )

    for task in scenario.tasks:
        if not _clean_text(task.title):
            errors.append(f"Задача {task.task_id}: название не может быть пустым.")
        historical_done = is_historical_done(
            effective_task_rows.get(task.task_id, {"status": task.status})
        )
        if historical_done:
            continue
        if task.team_id is not None and task.team_id not in known_teams:
            errors.append(
                f"Задача {task.task_id}: команда {task.team_id} "
                "не добавлена в сценарий."
            )
        role_names = [_clean_text(item.role) for item in task.role_requirements]
        if not role_names:
            errors.append(f"Задача {task.task_id}: укажите хотя бы одну требуемую роль.")
        if any(not role for role in role_names):
            errors.append(f"Задача {task.task_id}: роль не может быть пустой.")
        for duplicate in sorted(_duplicates(role_names)):
            errors.append(
                f"Задача {task.task_id}: роль {duplicate} указана "
                "более одного раза."
            )
        for dependency_id in task.depends_on:
            if dependency_id == task.task_id:
                errors.append(f"Задача {task.task_id} не может зависеть сама от себя.")
            elif dependency_is_satisfied(
                dependency_id,
                known_tasks=effective_task_rows,
                active_scenario_task_ids=active_scenario_task_ids,
            ):
                continue
            elif dependency_id not in catalog_task_rows:
                errors.append(
                    f"Задача {task.task_id}: блокирующая задача "
                    f"{dependency_id} неизвестна общему каталогу; не найдена."
                )
            else:
                errors.append(
                    f"Задача {task.task_id}: активная блокирующая задача "
                    f"{dependency_id} не включена в сценарий планирования."
                )

    cycle = _cycle_nodes(scenario.tasks, effective_task_rows)
    if cycle:
        errors.append(f"Обнаружен цикл зависимостей: {' → '.join(cycle)}.")
    if errors:
        raise ScenarioValidationError(errors)


def scenario_tables(
    scenario: ConstructorScenario,
    catalog_tables: dict[str, list[dict[str, Any]]] | None = None,
) -> dict[str, list[dict[str, Any]]]:
    """Convert a validated browser draft into normalized in-memory tables."""

    validate_constructor_scenario(scenario, catalog_tables)
    tasks: list[dict[str, Any]] = []
    role_effort: list[dict[str, Any]] = []
    dependencies: list[dict[str, Any]] = []
    engineers: list[dict[str, Any]] = []
    engineer_skills: list[dict[str, Any]] = []
    team_history: list[dict[str, Any]] = []
    task_skill_requirements: list[dict[str, Any]] = []
    selected_task_ids = {task.task_id for task in scenario.tasks}
    catalog_task_rows = {
        _clean_text(row.get("task_id")): row
        for row in (catalog_tables or {}).get("tasks", [])
        if _clean_text(row.get("task_id"))
    }
    # A browser draft contains work selected for the upcoming quarter, not the
    # complete planning context.  Every catalogued Done task is historical
    # context: it may satisfy a dependency and must be represented in the
    # draft summary, but it can never become work in a sprint.  Keeping the
    # whole set here makes a manual "select all" build equivalent to the
    # matching Excel/catalog build, including historical totals.
    historical_done_ids = {
        task_id
        for task_id, row in catalog_task_rows.items()
        if task_id not in selected_task_ids and is_historical_done(row)
    }

    for team in scenario.teams:
        team_id = _clean_text(team.team_id)
        history_row = {
            "team_id": team_id,
            "team_name": _clean_text(team.name) or team_id,
        }
        if team.expected_velocity_sp is not None:
            history_row["velocity_achieved"] = team.expected_velocity_sp
        if team.capacity_sp is not None:
            history_row["capacity_sp"] = team.capacity_sp
        team_history.append(history_row)

    for engineer in scenario.engineers:
        engineer_id = _clean_text(engineer.engineer_id)
        # Keep both team memberships and the explicit shared-pool assignment.
        # The browser sends an unassigned employee as ``team_id: null``; if we
        # only serialized team memberships here that employee disappeared from
        # the temporary planner context before feasibility was evaluated.
        for assignment in engineer.assignments:
            engineers.append(
                {
                    "engineer_id": engineer_id,
                    "name": _clean_text(engineer.name) or engineer_id,
                    "team_id": _clean_text(assignment.team_id) or None,
                    "role_raw": _clean_text(assignment.role),
                    "role": _clean_text(assignment.role),
                    "grade": _clean_text(assignment.grade) or None,
                    "capacity_rate": assignment.capacity_rate,
                }
            )
        for skill in sorted({canonicalize_skill(skill) for skill in engineer.skills if canonicalize_skill(skill)}):
            engineer_skills.append({"engineer_id": engineer_id, "technology": skill})

    for task in scenario.tasks:
        task_id = _clean_text(task.task_id)
        source_row = catalog_task_rows.get(task_id, {})
        effective_status = (
            "Done"
            if task.status == "Done" or is_historical_done(source_row)
            else task.status
        )
        tasks.append(
            {
                "task_id": task_id,
                "summary": _clean_text(task.title),
                "initiative_id": _clean_text(task.initiative_id) or None,
                # No source team means the planner evaluates every selected team
                # and chooses the earliest feasible assignment deterministically.
                "team_id": _clean_text(task.team_id) or None,
                "status": effective_status,
                "rung": task.rung,
                "story_points": task.story_points,
                "planned_start": _clean_text(task.planned_start) or None,
                "deadline": _clean_text(task.deadline) or None,
                "is_placeholder": False,
            }
        )
        explicit_skill_rows: list[dict[str, Any]] = [
            requirement.model_dump()
            for requirement in task.required_skills
            if not requirement.rejected
        ]
        rejected_skills = [
            requirement.technology
            for requirement in task.required_skills
            if requirement.rejected
        ]
        for requirement in task.role_requirements:
            role = _clean_text(requirement.role)
            role_effort.append(
                {
                    "task_id": task_id,
                    "role_raw": role,
                    "role": role,
                    "required_hours": requirement.required_hours,
                }
            )
            explicit_skill_rows.extend(
                {"technology": skill, "role": role, "source": "explicit"}
                for skill in requirement.skills
                if canonicalize_skill(skill)
            )
        task_skill_requirements.extend(
            merge_task_skill_requirements(
                task_id=task_id,
                title=task.title,
                explicit=explicit_skill_rows,
                roles=[requirement.role for requirement in task.role_requirements],
                rejected=rejected_skills,
            )
        )
        for technology in sorted({canonicalize_skill(value) for value in rejected_skills if canonicalize_skill(value)}):
            task_skill_requirements.append({
                "task_id": task_id,
                "role": None,
                "technology": technology,
                "source": "title",
                "matched_text": technology,
                "confidence": "high",
                "required_hours": None,
                "extractor_version": EXTRACTOR_VERSION,
                "confirmed": False,
                "rejected": True,
            })
        for blocking_task_id in sorted(set(task.depends_on)):
            dependencies.append(
                {
                    "blocking_task_id": blocking_task_id,
                    "blocked_task_id": task_id,
                    "dependency_type": "depends on",
                }
            )

    # Done-before-quarter tasks are planning facts, not selected work. They are
    # present in the preview as historical context and removed from Current
    # Plan only when a draft is approved.
    for task_id in sorted(historical_done_ids):
        source = catalog_task_rows[task_id]
        tasks.append(
            {
                "task_id": task_id,
                "summary": _clean_text(source.get("summary")) or "Unknown task",
                "initiative_id": _clean_text(source.get("initiative_id")) or task_id,
                "team_id": _clean_text(source.get("team_id")) or None,
                "status": "Done",
                "rung": source.get("rung"),
                "story_points": source.get("story_points"),
                "is_placeholder": False,
            }
        )

    return {
        "tasks": tasks,
        "task_role_effort": role_effort,
        "task_dependencies": dependencies,
        "engineers": engineers,
        "engineer_skills": engineer_skills,
        "team_history": team_history,
        "task_skill_requirements": task_skill_requirements,
    }


def _without_support_tasks(
    dashboard: dict[str, Any], support_task_ids: Iterable[str]
) -> dict[str, Any]:
    hidden = set(support_task_ids)
    if not hidden:
        return dashboard
    payload = deepcopy(dashboard)
    payload["tasks"] = [
        task for task in payload.get("tasks", []) if task.get("task_id") not in hidden
    ]
    plan = payload.get("plan", {})
    plan["tasks"] = [
        task for task in plan.get("tasks", []) if task.get("task_id") not in hidden
    ]
    plan["allocations"] = [
        row for row in plan.get("allocations", []) if row.get("task_id") not in hidden
    ]
    for summary in (payload.get("summary"), plan.get("summary")):
        if isinstance(summary, dict):
            summary["total_tasks"] = len(payload["tasks"])
    return payload


def _include_archived_done_blocker_context(
    tables: dict[str, list[dict[str, Any]]], repository: CatalogRepository
) -> dict[str, list[dict[str, Any]]]:
    """Keep archived Done blockers in planner context without reviving work."""

    result = deepcopy(tables)
    active_task_ids = {
        _clean_text(row.get("task_id")) for row in result.get("tasks", [])
    }
    archived_ids = repository.archived_ids("task")
    source_tables = repository.tables()
    source_tasks = {
        _clean_text(row.get("task_id")): row
        for row in source_tables.get("tasks", [])
        if _clean_text(row.get("task_id"))
    }
    referenced_done_ids = {
        blocker_id
        for dependency in result.get("task_dependencies", [])
        if _clean_text(dependency.get("blocked_task_id")) in active_task_ids
        and (blocker_id := _clean_text(dependency.get("blocking_task_id"))) in archived_ids
        and is_historical_done(source_tasks.get(blocker_id, {}))
    }
    present_ids = {
        _clean_text(row.get("task_id")) for row in result.get("tasks", [])
    }
    result.setdefault("tasks", []).extend(
        dict(source_tasks[task_id])
        for task_id in sorted(referenced_done_ids - present_ids)
    )
    return result


def build_constructor_plan(
    scenario: ConstructorScenario,
    catalog_tables: dict[str, list[dict[str, Any]]] | None = None,
) -> dict[str, Any]:
    """Run the real quarter planner against an ephemeral constructor scenario."""

    selected_ids = {task.task_id for task in scenario.tasks}
    tables = scenario_tables(scenario, catalog_tables)
    # Historical Done rows belong to the Base Plan summary.  There are no
    # non-historical support rows in a manual scenario, but retain the filter
    # for any future support-only entity type.
    support_ids = {
        _clean_text(row.get("task_id"))
        for row in tables["tasks"]
        if _clean_text(row.get("task_id")) not in selected_ids
        and not is_historical_done(row)
    }
    service = PlanningAPIService.from_tables(
        tables,
        plan_version="constructor-draft",
        planner_options={
            "allow_team_reassignment": False,
            # Team-bound cross-team help remains opt-in, while the planner
            # always admits unassigned/shared-pool engineers as helpers.
            "allow_cross_team": False,
        },
    )
    return {
        **_without_support_tasks(service.dashboard, support_ids),
        "_recommendation_context": {
            "tables": tables,
            "support_task_ids": sorted(support_ids),
        },
        "draft": {
            "storage": "browser_local_storage",
            "server_persisted": False,
            "source_dataset_modified": False,
            "team_count": len(scenario.teams),
            "engineer_count": len(scenario.engineers),
            "task_count": len(scenario.tasks),
        },
    }


def prepare_constructor_workbook(
    filename: str,
    content: bytes,
    repository: CatalogRepository,
) -> PreparedConstructorUpload:
    """Parse, normalize and audit an upload without merge or planner calls."""

    safe_name = Path(filename).name
    suffix = Path(safe_name).suffix.lower()
    if not safe_name or suffix not in {".xlsx", ".xlsm"}:
        raise ConstructorDatasetUploadError(
            "Поддерживаются только Excel-файлы .xlsx и .xlsm."
        )
    if not content:
        raise ConstructorDatasetUploadError("Загруженный Excel-файл пуст.")
    if len(content) > MAX_CONSTRUCTOR_UPLOAD_BYTES:
        raise ConstructorDatasetUploadError(
            "Размер Excel-файла не должен превышать 25 МБ."
        )

    try:
        with tempfile.TemporaryDirectory(prefix="plan-constructor-upload-") as directory:
            workbook = Path(directory) / safe_name
            workbook.write_bytes(content)
            parsed = parse_workbook(workbook)
            normalized, deduplication = _deduplicate_normalized_dataset(normalize(parsed))
            validation = validate(normalized)
            imported_tables = {
                table_name: normalized.tables[csv_name]
                for table_name, csv_name in TABLE_FILES.items()
            }
            for table_name, csv_name in {
                "task_skill_requirements": "task_skill_requirements.csv",
                "planning_results": "planning_results.csv",
                "plan_mismatch_reasons": "plan_mismatch_reasons.csv",
                "task_closure_results": "task_closure_results.csv",
            }.items():
                if csv_name in normalized.tables:
                    imported_tables[table_name] = normalized.tables[csv_name]
            fingerprint = hashlib.sha256(content).hexdigest()[:8]
            merge_preview = repository.merge_tables(imported_tables, dry_run=True)
            audit = build_data_preflight(
                parsed, normalized, validation, merge_preview
            )
            audit["deduplication"] = deduplication
    except (
        DatasetParsingError,
        BadZipFile,
        InvalidFileException,
        OSError,
        ValueError,
        KeyError,
    ) as error:
        raise ConstructorDatasetUploadError(str(error)) from error

    return PreparedConstructorUpload(
        filename=safe_name,
        size_bytes=len(content),
        fingerprint=fingerprint,
        tables=imported_tables,
        validation_warnings=tuple(validation.warnings),
        audit=audit,
    )


def build_constructor_plan_from_prepared(
    prepared: PreparedConstructorUpload,
    repository: CatalogRepository,
    archive_decisions: dict[str, str] | None = None,
    existing_decisions: dict[str, str] | None = None,
) -> dict[str, Any]:
    """Merge an accepted audited snapshot and build its Base Plan once."""

    if prepared.audit.get("has_critical_errors"):
        raise ConstructorDatasetUploadError(
            "В данных есть ошибки, из-за которых план невозможно построить."
        )

    try:
        decisions = _validated_archive_decisions(prepared, archive_decisions)
        active_decisions = _validated_existing_decisions(prepared, existing_decisions)
        merge_result = repository.merge_tables(
            prepared.tables,
            archive_decisions=decisions,
            existing_decisions=active_decisions,
        )
        planning_tables = _include_archived_done_blocker_context(
            repository.active_planning_tables(), repository
        )
        service = PlanningAPIService.from_tables(
            planning_tables,
            plan_version=f"constructor-upload-{prepared.fingerprint}",
            planner_options={
                "allow_team_reassignment": False,
                "allow_cross_team": False,
            },
        )
    except (ServiceDataError, OSError, ValueError, KeyError) as error:
        raise ConstructorDatasetUploadError(str(error)) from error

    hidden_archived_tasks = repository.archived_ids("task")
    catalog = build_constructor_catalog(
        service, archived_task_ids=hidden_archived_tasks
    )
    return {
        **_without_support_tasks(service.dashboard, hidden_archived_tasks),
        "_recommendation_context": {
            "tables": planning_tables,
            "support_task_ids": sorted(hidden_archived_tasks),
        },
        "catalog": catalog,
        "uploaded_dataset": {
            "name": prepared.filename,
            "size_bytes": prepared.size_bytes,
            "fingerprint": prepared.fingerprint,
            "warnings": [*prepared.validation_warnings, *merge_result.warnings],
            "catalog_merge": merge_result.to_dict(),
            "server_persisted": True,
            "source_dataset_modified": False,
            "data_preflight": prepared.audit,
        },
    }


def build_constructor_plan_from_workbook(
    filename: str,
    content: bytes,
    repository: CatalogRepository,
) -> dict[str, Any]:
    """Compatibility wrapper for trusted callers that need immediate build."""

    prepared = prepare_constructor_workbook(filename, content, repository)
    return build_constructor_plan_from_prepared(prepared, repository)


def build_constructor_catalog(
    service: PlanningAPIService,
    *,
    archived_task_ids: Iterable[str] = (),
) -> dict[str, Any]:
    """Expose active catalog entities plus preserved dependency references."""

    archived_tasks = set(archived_task_ids)

    # Keep historical velocity and manual capacity separate. A new team has
    # no velocity samples yet, so its capacity must not be exposed as a fake
    # zero velocity when the catalog is reopened.
    historical_velocity_by_team: dict[str, list[float]] = defaultdict(list)
    manual_capacity_by_team: dict[str, list[float]] = defaultdict(list)
    for history_row in service.tables.get("team_history", []):
        history_team_id = _clean_text(history_row.get("team_id"))
        if not history_team_id:
            continue
        try:
            velocity = float(history_row.get("velocity_achieved"))
            if velocity >= 0:
                historical_velocity_by_team[history_team_id].append(velocity)
        except (TypeError, ValueError):
            pass
        try:
            capacity = float(history_row.get("capacity_sp"))
            if capacity >= 0:
                manual_capacity_by_team[history_team_id].append(capacity)
        except (TypeError, ValueError):
            pass

    skills_by_engineer: dict[str, set[str]] = defaultdict(set)
    for row in service.tables["engineer_skills"]:
        engineer_id = _clean_text(row.get("engineer_id"))
        skill = _clean_text(row.get("technology"))
        if engineer_id and skill:
            skills_by_engineer[engineer_id].add(skill)

    engineer_profiles: dict[str, dict[str, Any]] = {}
    seen_assignments: set[tuple[str, str, str]] = set()
    for row in service.tables["engineers"]:
        engineer_id = _clean_text(row.get("engineer_id"))
        team_id = _clean_text(row.get("team_id"))
        role = _clean_text(row.get("role"))
        if not engineer_id or not role:
            continue
        target = engineer_profiles.setdefault(
            engineer_id,
            {
                "engineer_id": engineer_id,
                "name": _clean_text(row.get("name")) or engineer_id,
                "assignments": [],
                "free_assignment": None,
                "skills": [],
                "source": "dataset",
            },
        )
        if not team_id:
            if target["free_assignment"] is None:
                target["free_assignment"] = {
                    "role": role,
                    "grade": _clean_text(row.get("grade")) or None,
                    "capacity_rate": float(row.get("capacity_rate") or 1),
                }
            continue
        key = (engineer_id, team_id, role)
        if key not in seen_assignments:
            target["assignments"].append(
                {
                    "team_id": team_id,
                    "role": role,
                    "grade": _clean_text(row.get("grade")) or None,
                    "capacity_rate": float(row.get("capacity_rate") or 0),
                }
            )
            seen_assignments.add(key)
    for engineer_id, profile in engineer_profiles.items():
        profile["assignments"].sort(key=lambda item: (item["team_id"], item["role"]))
        profile["is_free"] = not profile["assignments"] and bool(
            profile["free_assignment"]
        )
        profile["skills"] = sorted(skills_by_engineer[engineer_id])

    role_requirements: dict[str, list[dict[str, Any]]] = defaultdict(list)
    skill_requirements: dict[tuple[str, str], list[str]] = defaultdict(list)
    task_skill_details: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in service.tables.get("task_skill_requirements", []):
        task_id = _clean_text(row.get("task_id"))
        role = _clean_text(row.get("role")) or None
        key = (task_id, role or "")
        technology = canonicalize_skill(row.get("technology"))
        if task_id and technology:
            task_skill_details[task_id].append({
                "technology": technology,
                "role": role,
                "source": _clean_text(row.get("source")) or "explicit",
                "confidence": _clean_text(row.get("confidence")) or "high",
                "matched_text": _clean_text(row.get("matched_text")) or None,
                "confirmed": str(row.get("confirmed", "true")).strip().casefold() in {"true", "1", "yes"},
                "rejected": str(row.get("rejected", "false")).strip().casefold() in {"true", "1", "yes"},
            })
        if task_id and role and technology and str(row.get("rejected", "false")).strip().casefold() not in {"true", "1", "yes"}:
            skill_requirements[key].append(technology)
    for row in service.tables["task_role_effort"]:
        task_id = _clean_text(row.get("task_id"))
        role = _clean_text(row.get("role"))
        hours = row.get("required_hours")
        try:
            required_hours = float(hours)
        except (TypeError, ValueError):
            continue
        if task_id and role and required_hours > 0:
            role_requirements[task_id].append(
                {
                    "role": role,
                    "required_hours": required_hours,
                    "skills": sorted(set(skill_requirements[(task_id, role)])),
                }
            )

    dependencies: dict[str, list[str]] = defaultdict(list)
    for row in service.tables["task_dependencies"]:
        blocker = _clean_text(row.get("blocking_task_id"))
        blocked = _clean_text(row.get("blocked_task_id"))
        if blocker and blocked:
            dependencies[blocked].append(blocker)

    raw_tasks = {
        _clean_text(row.get("task_id")): row
        for row in service.tables["tasks"]
        if _clean_text(row.get("task_id"))
    }
    catalog_tasks = []
    for task in service.tasks:
        if task["is_placeholder"] or task["task_id"] in archived_tasks:
            continue
        task_id = task["task_id"]
        raw = raw_tasks[task_id]
        team_id = _clean_text(raw.get("team_id")) or task["original_team_id"]
        status = _clean_text(raw.get("status")) or "ToDo"
        if status.casefold() not in {"todo", "inprogress", "done"}:
            status = "ToDo"
        status = {"todo": "ToDo", "inprogress": "InProgress", "done": "Done"}[
            status.casefold()
        ]
        if status != "Done" and (not team_id or not role_requirements[task_id]):
            continue
        catalog_tasks.append(
            {
                "task_id": task_id,
                "title": task["title"],
                "initiative_id": task["initiative_id"],
                "team_id": team_id or None,
                "status": status,
                "rung": task["rung"] or 0,
                "story_points": task["story_points"] or 0,
                "planned_start": _clean_text(raw.get("planned_start")) or None,
                "deadline": _clean_text(raw.get("deadline")) or None,
                "role_requirements": sorted(
                    role_requirements[task_id], key=lambda item: item["role"]
                ),
                "required_skills": sorted(
                    task_skill_details[task_id],
                    key=lambda item: ((item.get("role") or ""), item["technology"].casefold()),
                ),
                "depends_on": sorted(set(dependencies[task_id])),
                "source": "dataset",
            }
        )

    known_task_ids = set(raw_tasks)
    for task in catalog_tasks:
        task["depends_on"] = [
            task_id for task_id in task["depends_on"] if task_id in known_task_ids
        ]

    roles = sorted(
        {
            assignment["role"]
            for profile in engineer_profiles.values()
            for assignment in profile["assignments"]
        }
        | {
            item["role"]
            for requirements in role_requirements.values()
            for item in requirements
        }
    )
    skill_registry = get_skill_registry()
    skills = [item.display_name for item in skill_registry.definitions]
    grades = sorted(
        {
            assignment["grade"]
            for profile in engineer_profiles.values()
            for assignment in profile["assignments"]
            if assignment["grade"]
        }
    )
    initialized = bool(
        catalog_tasks
        or engineer_profiles
        or service.teams
        or any(service.tables.get(name) for name in ("tasks", "engineers", "team_history"))
    )
    return {
        "status": "READY" if initialized else "EMPTY",
        "data_state": "READY" if initialized else "UNINITIALIZED",
        "message": None if initialized else "Каталог ещё не заполнен",
        "teams": [
            {
                "team_id": team["team_id"],
                "name": team["name"],
                "expected_velocity_sp": (
                    sum(historical_velocity_by_team[team["team_id"]])
                    / len(historical_velocity_by_team[team["team_id"]])
                    if historical_velocity_by_team[team["team_id"]]
                    else None
                ),
                "capacity_sp": (
                    sum(manual_capacity_by_team[team["team_id"]])
                    / len(manual_capacity_by_team[team["team_id"]])
                    if manual_capacity_by_team[team["team_id"]]
                    else None
                ),
                "focus_factor": 0.8,
                "source": "dataset",
            }
            for team in service.teams
        ],
        "engineers": sorted(engineer_profiles.values(), key=lambda item: item["engineer_id"]),
        "tasks": sorted(catalog_tasks, key=lambda item: item["task_id"]),
        "options": {
            "roles": roles,
            "skills": skills,
            "skill_registry": skill_registry.options(),
            "grades": grades,
            "statuses": ["ToDo", "InProgress", "Done"],
            "focus_factor": 0.8,
        },
        "source_dataset_read_only": True,
        "working_catalog_persistent": True,
    }


def build_catalog_snapshot(
    tables: dict[str, list[dict[str, Any]]],
    *,
    archived_task_ids: Iterable[str] = (),
    revision: int = 0,
) -> dict[str, Any]:
    """Build the catalog DTO without invoking the quarter planner.

    Catalog pages and mutations need current master data, not a simulated plan.
    Keeping this projection table-only also makes an engineer/team/task create
    cheap and prevents catalog writes from changing Approved Plan indirectly.
    """

    archived_tasks = set(archived_task_ids)

    skills_by_engineer: dict[str, set[str]] = defaultdict(set)
    for row in tables.get("engineer_skills", []):
        engineer_id = _clean_text(row.get("engineer_id"))
        skill = canonicalize_skill(row.get("technology"))
        if engineer_id and skill:
            skills_by_engineer[engineer_id].add(skill)

    profiles: dict[str, dict[str, Any]] = {}
    for row in tables.get("engineers", []):
        engineer_id = _clean_text(row.get("engineer_id"))
        role = _clean_text(row.get("role"))
        if not engineer_id or not role:
            continue
        profile = profiles.setdefault(engineer_id, {
            "engineer_id": engineer_id,
            "name": _clean_text(row.get("name")) or engineer_id,
            "assignments": [],
            "free_assignment": None,
            "skills": [],
            "source": "catalog",
        })
        assignment = {
            "team_id": _clean_text(row.get("team_id")) or None,
            "role": role,
            "grade": _clean_text(row.get("grade")) or None,
            "capacity_rate": float(row.get("capacity_rate") or 0),
        }
        if assignment["team_id"]:
            profile["assignments"].append(assignment)
        elif profile["free_assignment"] is None:
            profile["free_assignment"] = {key: value for key, value in assignment.items() if key != "team_id"}
    for engineer_id, profile in profiles.items():
        profile["assignments"].sort(key=lambda item: (item["team_id"], item["role"]))
        profile["skills"] = sorted(skills_by_engineer[engineer_id], key=str.casefold)
        profile["is_free"] = not profile["assignments"]
        profile["membership_rate"] = sum(float(item["capacity_rate"]) for item in profile["assignments"])
        profile["available_capacity_rate"] = (
            float(profile["free_assignment"].get("capacity_rate") or 0)
            if profile["free_assignment"]
            else max(1.0 - profile["membership_rate"], 0.0)
        )

    role_effort: dict[str, list[dict[str, Any]]] = defaultdict(list)
    skills_by_task_role: dict[tuple[str, str], set[str]] = defaultdict(set)
    task_skill_details: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in tables.get("task_skill_requirements", []):
        task_id = _clean_text(row.get("task_id"))
        role = _clean_text(row.get("role"))
        technology = canonicalize_skill(row.get("technology"))
        if not task_id or not technology:
            continue
        detail = {
            "technology": technology,
            "role": role or None,
            "source": _clean_text(row.get("source")) or "explicit",
            "confidence": _clean_text(row.get("confidence")) or "high",
            "matched_text": _clean_text(row.get("matched_text")) or None,
            "confirmed": str(row.get("confirmed", "true")).casefold() in {"true", "1", "yes"},
            "rejected": str(row.get("rejected", "false")).casefold() in {"true", "1", "yes"},
        }
        task_skill_details[task_id].append(detail)
        if role and not detail["rejected"]:
            skills_by_task_role[(task_id, role)].add(technology)
    for row in tables.get("task_role_effort", []):
        task_id = _clean_text(row.get("task_id"))
        role = _clean_text(row.get("role"))
        try:
            hours = float(row.get("required_hours") or 0)
        except (TypeError, ValueError):
            continue
        if task_id and role and hours > 0:
            role_effort[task_id].append({
                "role": role,
                "required_hours": hours,
                "skills": sorted(skills_by_task_role[(task_id, role)], key=str.casefold),
            })
    dependencies: dict[str, list[str]] = defaultdict(list)
    for row in tables.get("task_dependencies", []):
        blocker = _clean_text(row.get("blocking_task_id"))
        blocked = _clean_text(row.get("blocked_task_id"))
        if blocker and blocked:
            dependencies[blocked].append(blocker)

    tasks = []
    for row in tables.get("tasks", []):
        task_id = _clean_text(row.get("task_id"))
        if not task_id or task_id in archived_tasks or bool(row.get("is_placeholder")):
            continue
        raw_status = _clean_text(row.get("status")).casefold()
        status = {"todo": "ToDo", "inprogress": "InProgress", "done": "Done"}.get(raw_status, "ToDo")
        tasks.append({
            "task_id": task_id,
            "title": _clean_text(row.get("summary")) or task_id,
            "initiative_id": _clean_text(row.get("initiative_id")) or None,
            "team_id": _clean_text(row.get("team_id")) or None,
            "status": status,
            "rung": float(row.get("rung") or 0),
            "story_points": float(row.get("story_points") or 0),
            "planned_start": _clean_text(row.get("planned_start")) or None,
            "deadline": _clean_text(row.get("deadline")) or None,
            "role_requirements": sorted(role_effort[task_id], key=lambda item: item["role"]),
            "required_skills": sorted(task_skill_details[task_id], key=lambda item: ((item.get("role") or ""), item["technology"].casefold())),
            "depends_on": sorted(set(dependencies[task_id])),
            "source": "catalog",
        })

    team_names: dict[str, str] = {}
    team_velocity: dict[str, list[float]] = defaultdict(list)
    team_capacity: dict[str, list[float]] = defaultdict(list)
    for row in tables.get("team_history", []):
        team_id = _clean_text(row.get("team_id"))
        if not team_id:
            continue
        team_names[team_id] = _clean_text(row.get("team_name") or row.get("name")) or team_id
        value = row.get("velocity_achieved")
        try:
            if value is not None:
                team_velocity[team_id].append(float(value))
        except (TypeError, ValueError):
            pass
        manual = row.get("capacity_sp")
        try:
            if manual is not None:
                team_capacity[team_id].append(float(manual))
        except (TypeError, ValueError):
            pass
    team_ids = set(team_names)
    team_ids.update(
        _clean_text(item.get("team_id"))
        for item in tables.get("engineers", [])
        if _clean_text(item.get("team_id"))
    )
    team_ids.update(task["team_id"] for task in tasks if task["team_id"])
    teams = [{
        "team_id": team_id,
        "name": team_names.get(team_id, team_id),
        "expected_velocity_sp": (
            sum(team_velocity[team_id]) / len(team_velocity[team_id])
            if team_velocity[team_id] else None
        ),
        "capacity_sp": (
            sum(team_capacity[team_id]) / len(team_capacity[team_id])
            if team_capacity[team_id] else None
        ),
        "focus_factor": 0.8,
        "source": "catalog",
    } for team_id in sorted(team_ids)]

    roles = sorted({
        item["role"]
        for profile in profiles.values()
        for item in [*profile["assignments"], *([profile["free_assignment"]] if profile["free_assignment"] else [])]
        if item.get("role")
    } | {
        item["role"] for rows in role_effort.values() for item in rows
    }, key=str.casefold)
    grades = sorted({
        item["grade"]
        for profile in profiles.values()
        for item in [*profile["assignments"], *([profile["free_assignment"]] if profile["free_assignment"] else [])]
        if item.get("grade")
    }, key=str.casefold)
    registry = get_skill_registry()
    registry_skills = {item.display_name for item in registry.definitions}
    all_skills = registry_skills | {
        skill for values in skills_by_engineer.values() for skill in values
    } | {
        item["technology"] for values in task_skill_details.values() for item in values
    }
    initialized = bool(tasks or profiles or teams)
    return {
        "status": "READY" if initialized else "EMPTY",
        "data_state": "READY" if initialized else "UNINITIALIZED",
        "message": None if initialized else "Каталог ещё не заполнен",
        "catalog_revision": revision,
        "teams": teams,
        "engineers": sorted(profiles.values(), key=lambda item: item["engineer_id"]),
        "tasks": sorted(tasks, key=lambda item: item["task_id"]),
        "options": {
            "roles": roles,
            "skills": sorted(all_skills, key=str.casefold),
            "skill_registry": registry.options(),
            "grades": grades,
            "statuses": ["ToDo", "InProgress", "Done"],
            "focus_factor": 0.8,
        },
        "source_dataset_read_only": True,
        "working_catalog_persistent": True,
    }
