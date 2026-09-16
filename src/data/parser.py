"""Detect and extract logical tables from the organizer's single Excel sheet."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


class DatasetParsingError(RuntimeError):
    """A required table could not be located or decoded from the raw workbook."""


TASK_ID_PATTERN = re.compile(r"^[A-Za-zА-Яа-я]+-\d+[A-Za-z0-9-]*$")


@dataclass(frozen=True)
class RawTable:
    name: str
    marker_row: int | None
    header_row: int
    headers: list[Any]
    rows: list[list[Any]]


@dataclass(frozen=True)
class RawReferenceList:
    """A titled, single-column reference list from the source worksheet."""

    name: str
    marker_row: int
    title: str
    values: list[Any]


@dataclass(frozen=True)
class ParsedDataset:
    tasks: RawTable
    effort: RawTable
    dependencies: RawTable
    engineers: RawTable
    team_history: RawTable
    planning_results: RawReferenceList
    plan_mismatch_reasons: RawReferenceList
    task_closure_results: RawReferenceList


REFERENCE_LIST_TITLES = {
    "planning_results": (
        "Итог разработки, выбирается заказчиком и исполнителем на этапе планирования"
    ),
    "plan_mismatch_reasons": "Причины расхождения планов заказчика и исполнителя",
    "task_closure_results": "Результаты работ при закрытии задачи",
}


def normalized_text(value: Any) -> str:
    """Return a whitespace-normalized text token for header and marker matching."""

    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return " ".join(str(value).split()).casefold()


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def row_values(sheet: Worksheet, row_index: int) -> list[Any]:
    values = [sheet.cell(row_index, column).value for column in range(1, sheet.max_column + 1)]
    last_nonblank = 0
    for index, value in enumerate(values, start=1):
        if not is_blank(value):
            last_nonblank = index
    return values[:last_nonblank]


def row_has_values(values: Iterable[Any]) -> bool:
    return any(not is_blank(value) for value in values)


def find_header_row(sheet: Worksheet, required_headers: set[str]) -> int:
    expected = {normalized_text(header) for header in required_headers}
    for row_index in range(1, sheet.max_row + 1):
        row_headers = {normalized_text(value) for value in row_values(sheet, row_index)}
        if expected.issubset(row_headers):
            return row_index
    readable_headers = ", ".join(sorted(required_headers))
    raise DatasetParsingError(f"Could not find a header row containing: {readable_headers}.")


def preceding_marker_row(sheet: Worksheet, header_row: int) -> int | None:
    """Return the nearest non-empty row before a header, if it is a title marker."""

    for row_index in range(header_row - 1, 0, -1):
        values = row_values(sheet, row_index)
        if row_has_values(values):
            return row_index if len([value for value in values if not is_blank(value)]) == 1 else None
    return None


def contiguous_rows(
    sheet: Worksheet,
    header_row: int,
    width: int,
    key_column: int = 0,
    stop_tokens: set[str] | None = None,
) -> list[list[Any]]:
    """Read records below a header until a blank row or a logical total marker."""

    rows: list[list[Any]] = []
    stop_tokens = stop_tokens or set()
    for row_index in range(header_row + 1, sheet.max_row + 1):
        values = row_values(sheet, row_index)[:width]
        if not row_has_values(values):
            break
        first_token = normalized_text(values[0] if values else None)
        if first_token in stop_tokens or any(first_token.startswith(token) for token in stop_tokens):
            break
        if key_column >= len(values) or is_blank(values[key_column]):
            break
        rows.append(values + [None] * (width - len(values)))
    return rows


def find_effort_header_row(sheet: Worksheet) -> tuple[int, int | None]:
    """Find the role/task wide matrix using its title and the Role column marker."""

    marker_row: int | None = None
    for row_index in range(1, sheet.max_row + 1):
        values = row_values(sheet, row_index)
        first_cell = normalized_text(values[0] if values else None)
        if "estimates" in first_cell or "калькулятор сметы" in first_cell:
            marker_row = row_index
            continue
        if marker_row is not None and first_cell in {"роль", "role"} and len(values) > 1:
            return row_index, marker_row
    raise DatasetParsingError("Could not find the Estimates role/task matrix.")


def looks_like_task_id(value: Any) -> bool:
    return bool(TASK_ID_PATTERN.fullmatch(str(value).strip())) if value is not None else False


def parse_dependency_values(values: list[Any]) -> list[Any] | None:
    """Read the three raw dependency values despite merged/misaligned headers.

    The supplied sheet places its header cells one column to the right of its data
    cells. Data rows themselves are unambiguous: blocking id, blocked id, relation
    text. No dependency type is interpreted and no direction is reversed here.
    """

    nonblank = [value for value in values if not is_blank(value)]
    task_positions = [index for index, value in enumerate(nonblank) if looks_like_task_id(value)]
    if len(task_positions) < 2:
        return None
    predecessor = nonblank[task_positions[0]]
    successor = nonblank[task_positions[1]]
    dependency_type = next(
        (value for value in nonblank[task_positions[1] + 1 :] if not looks_like_task_id(value)),
        None,
    )
    return [predecessor, successor, dependency_type]


def extract_dependencies(sheet: Worksheet, header_row: int) -> list[list[Any]]:
    records: list[list[Any]] = []
    for row_index in range(header_row + 1, sheet.max_row + 1):
        values = row_values(sheet, row_index)
        if not row_has_values(values):
            break
        record = parse_dependency_values(values)
        if record is None:
            raise DatasetParsingError(
                f"Dependency row {row_index} does not contain two task identifiers."
            )
        records.append(record)
    return records


def find_reference_list(sheet: Worksheet, name: str, title: str) -> RawReferenceList:
    """Locate a titled one-column list and preserve its values in source order."""

    expected_title = normalized_text(title)
    marker_row: int | None = None
    for row_index in range(1, sheet.max_row + 1):
        nonblank = [value for value in row_values(sheet, row_index) if not is_blank(value)]
        if len(nonblank) == 1 and normalized_text(nonblank[0]) == expected_title:
            marker_row = row_index
            break
    if marker_row is None:
        raise DatasetParsingError(f"Could not find reference list: {title}.")

    values: list[Any] = []
    for row_index in range(marker_row + 1, sheet.max_row + 1):
        nonblank = [value for value in row_values(sheet, row_index) if not is_blank(value)]
        if not nonblank:
            break
        if len(nonblank) != 1:
            raise DatasetParsingError(
                f"Reference list {title!r} row {row_index} must contain exactly one value."
            )
        values.append(nonblank[0])
    if not values:
        raise DatasetParsingError(f"Reference list {title!r} is empty.")
    return RawReferenceList(name=name, marker_row=marker_row, title=title, values=values)


def detect_blocks(sheet: Worksheet) -> dict[str, RawTable | RawReferenceList]:
    """Locate every logical table from headers and nearby title markers."""

    tasks_header = find_header_row(sheet, {"task_id", "team_id", "status", "estimation_sp"})
    tasks_headers = row_values(sheet, tasks_header)

    effort_header, effort_marker = find_effort_header_row(sheet)
    effort_headers = row_values(sheet, effort_header)

    dependencies_header = find_header_row(
        sheet, {"blocking_task_id", "blocked_task_id", "dependency_type"}
    )
    engineers_header = find_header_row(
        sheet, {"engineer_id", "team_id", "role", "grade", "skills_declared", "capacity_rate"}
    )
    engineers_headers = row_values(sheet, engineers_header)
    history_header = find_header_row(
        sheet, {"snapshot_date", "team_id", "velocity_achieved", "planned_sp"}
    )
    history_headers = row_values(sheet, history_header)

    blocks: dict[str, RawTable | RawReferenceList] = {
        "tasks": RawTable(
            name="tasks",
            marker_row=preceding_marker_row(sheet, tasks_header),
            header_row=tasks_header,
            headers=tasks_headers,
            rows=contiguous_rows(sheet, tasks_header, len(tasks_headers), key_column=1),
        ),
        "effort": RawTable(
            name="effort",
            marker_row=effort_marker,
            header_row=effort_header,
            headers=effort_headers,
            rows=contiguous_rows(
                sheet,
                effort_header,
                len(effort_headers),
                stop_tokens={"итого", "total"},
            ),
        ),
        "dependencies": RawTable(
            name="dependencies",
            marker_row=preceding_marker_row(sheet, dependencies_header),
            header_row=dependencies_header,
            headers=["blocking_task_id", "blocked_task_id", "dependency_type"],
            rows=extract_dependencies(sheet, dependencies_header),
        ),
        "engineers": RawTable(
            name="engineers",
            marker_row=preceding_marker_row(sheet, engineers_header),
            header_row=engineers_header,
            headers=engineers_headers,
            rows=contiguous_rows(sheet, engineers_header, len(engineers_headers), key_column=0),
        ),
        "team_history": RawTable(
            name="team_history",
            marker_row=preceding_marker_row(sheet, history_header),
            header_row=history_header,
            headers=history_headers,
            rows=contiguous_rows(sheet, history_header, len(history_headers), key_column=0),
        ),
    }
    for name, title in REFERENCE_LIST_TITLES.items():
        blocks[name] = find_reference_list(sheet, name, title)
    return blocks


def parse_workbook(raw_file: Path) -> ParsedDataset:
    """Load the single organizer sheet with cached formula values, never modifying it."""

    try:
        workbook = load_workbook(raw_file, data_only=True, read_only=False)
    except Exception as error:  # openpyxl has several format-specific exception classes.
        raise DatasetParsingError(f"Could not open raw workbook {raw_file}: {error}") from error

    if not workbook.worksheets:
        raise DatasetParsingError(f"Raw workbook {raw_file} has no worksheets.")
    blocks = detect_blocks(workbook.worksheets[0])
    return ParsedDataset(
        tasks=blocks["tasks"],
        effort=blocks["effort"],
        dependencies=blocks["dependencies"],
        engineers=blocks["engineers"],
        team_history=blocks["team_history"],
        planning_results=blocks["planning_results"],
        plan_mismatch_reasons=blocks["plan_mismatch_reasons"],
        task_closure_results=blocks["task_closure_results"],
    )
