from __future__ import annotations

import csv
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts.prepare_dataset import run_preprocessing
from src.data.normalization import melt_effort, normalize, split_skills
from src.data.parser import RawTable, parse_workbook
from src.data.validation import dependency_reference_issues, validate, validate_capacity_rate


def build_fixture_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Tasks"
    sheet.append(
        [
            "Номер инициативы",
            "task_id",
            "summary",
            "team_id",
            "status",
            "estimation_sp",
            "rung",
            "planned_end",
        ]
    )
    sheet.append(["PRODF-1", "T-1", "Real task 1", "Team-A", "ToDo", 3, 10, "2026-03-31"])
    sheet.append(["PRODF-1", "T-2", "Real task 2", "Team-A", "ToDo", 5, 20, "2026-04-15"])
    sheet.append([])
    sheet.append(["Estimates. Калькулятор сметы"])
    sheet.append(["Роль", "T-1", "T-2", "T-3"])
    sheet.append(["Аналитик", 8, 0, 2])
    sheet.append(["Итого по ролям", 8, 0, 2])
    sheet.append([])
    sheet.append(["Task_dependencies"])
    # The organizer file has an offset dependency header, so the fixture has one too.
    sheet.append(["blocking_task_id", None, "blocked_task_id", "dependency_type"])
    sheet.append(["T-4", "T-2", "has to be done before"])
    sheet.append(["T-1", "T-4", "depends on"])
    sheet.append([])
    sheet.append(["Engineers_profiles"])
    sheet.append(["engineer_id", "team_id", "role", "grade", "skills_declared", "capacity_rate"])
    sheet.append(["ENG-1", "Team-A", "Аналитик", "Middle", "SQL, CI/CD (GitLab, Jenkins)", "1.0"])
    sheet.append([])
    sheet.append(["Team_history"])
    sheet.append(["snapshot_date", "team_id", "velocity_achieved", "planned_sp"])
    sheet.append(["2026-01-01", "Team-A", 8, 10])
    sheet.append([])
    sheet.append(
        ["Итог разработки, выбирается заказчиком и исполнителем на этапе планирования"]
    )
    for value in (
        "1. Аналитика и архитектура выполнены частично",
        "2. Аналитика и архитектура завершены",
        "3. Функционал реализован частично, работы будут продолжены",
        "4. Функционал реализован и готов к ИТ",
        "5. Функционал реализован и перенесен в опытную эксплуатацию",
        "6. Функционал реализован и перенесен в промышленную эксплуатацию",
        "не будет взято в квартал",
        "Отменено зака",
    ):
        sheet.append([value])
    sheet.append([])
    sheet.append(["Причины расхождения планов заказчика и исполнителя"])
    for value in (
        "Не оценивалась",
        "Отсутствие ресурсов",
        "Отсутствует готовность смежных команд",
        "Превышение плана",
        "Изменились требования",
    ):
        sheet.append([value])
    sheet.append([])
    sheet.append(["Результаты работ при закрытии задачи"])
    for value in (
        "планируемый результат работ достигнут",
        "планируемый результат работ НЕ достигнут",
        "Работы отменены/перенесены заказчиком",
    ):
        sheet.append([value])
    workbook.save(path)


def normalize_fixture() -> object:
    with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as temporary_directory:
        raw_file = Path(temporary_directory) / "fixture.xlsx"
        build_fixture_workbook(raw_file)
        return normalize(parse_workbook(raw_file))


class ParserTests(unittest.TestCase):
    def test_detects_blocks_by_headers_and_markers(self) -> None:
        with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as temporary_directory:
            raw_file = Path(temporary_directory) / "fixture.xlsx"
            build_fixture_workbook(raw_file)
            dataset = parse_workbook(raw_file)
        self.assertEqual(len(dataset.tasks.rows), 2)
        self.assertEqual(dataset.effort.marker_row, 6)
        self.assertEqual(len(dataset.dependencies.rows), 2)
        self.assertEqual(len(dataset.engineers.rows), 1)
        self.assertEqual(len(dataset.team_history.rows), 1)
        self.assertEqual(len(dataset.planning_results.values), 8)
        self.assertEqual(len(dataset.plan_mismatch_reasons.values), 5)
        self.assertEqual(len(dataset.task_closure_results.values), 3)
        self.assertEqual(dataset.planning_results.values[-1], "Отменено зака")


class NormalizationTests(unittest.TestCase):
    def test_reference_lists_are_normalized_to_separate_ordered_tables(self) -> None:
        normalized = normalize_fixture()
        self.assertEqual(
            normalized.tables["planning_results.csv"][-1],
            {"source_order": 8, "value": "Отменено зака"},
        )
        self.assertEqual(
            normalized.tables["plan_mismatch_reasons.csv"][1],
            {"source_order": 2, "value": "Отсутствие ресурсов"},
        )
        self.assertEqual(
            normalized.tables["task_closure_results.csv"][0],
            {"source_order": 1, "value": "планируемый результат работ достигнут"},
        )

    def test_effort_only_task_id_creates_placeholder(self) -> None:
        normalized = normalize_fixture()
        tasks_by_id = {task["task_id"]: task for task in normalized.tables["tasks.csv"]}
        self.assertIn("T-3", tasks_by_id)
        self.assertTrue(tasks_by_id["T-3"]["is_placeholder"])

    def test_dependency_only_task_id_creates_placeholder(self) -> None:
        normalized = normalize_fixture()
        tasks_by_id = {task["task_id"]: task for task in normalized.tables["tasks.csv"]}
        self.assertIn("T-4", tasks_by_id)
        self.assertTrue(tasks_by_id["T-4"]["is_placeholder"])

    def test_repeated_unknown_id_creates_one_placeholder(self) -> None:
        normalized = normalize_fixture()
        task_ids = [task["task_id"] for task in normalized.tables["tasks.csv"]]
        self.assertEqual(task_ids.count("T-4"), 1)

    def test_real_task_does_not_become_placeholder(self) -> None:
        normalized = normalize_fixture()
        tasks_by_id = {task["task_id"]: task for task in normalized.tables["tasks.csv"]}
        self.assertFalse(tasks_by_id["T-1"]["is_placeholder"])
        self.assertEqual(tasks_by_id["T-1"]["summary"], "Real task 1")

    def test_placeholder_flag_and_name_are_explicit(self) -> None:
        normalized = normalize_fixture()
        task = next(task for task in normalized.tables["tasks.csv"] if task["task_id"] == "T-3")
        self.assertIs(task["is_placeholder"], True)
        self.assertEqual(task["summary"], "Unknown task")

    def test_placeholder_business_fields_remain_empty(self) -> None:
        normalized = normalize_fixture()
        task = next(task for task in normalized.tables["tasks.csv"] if task["task_id"] == "T-4")
        for field in (
            "initiative_id",
            "team_id",
            "status",
            "story_points",
            "rung",
            "planned_end",
            "spent_time",
            "estimated_hh",
        ):
            self.assertIsNone(task[field], field)

    def test_dependency_with_placeholder_is_preserved_without_reversal(self) -> None:
        normalized = normalize_fixture()
        self.assertIn(
            {
                "blocking_task_id": "T-4",
                "blocked_task_id": "T-2",
                "dependency_type": "has to be done before",
            },
            normalized.tables["task_dependencies.csv"],
        )
        self.assertIn(
            {
                "blocking_task_id": "T-1",
                "blocked_task_id": "T-4",
                "dependency_type": "depends on",
            },
            normalized.tables["task_dependencies.csv"],
        )

    def test_reference_validation_passes_after_placeholders(self) -> None:
        normalized = normalize_fixture()
        report = validate(normalized)
        self.assertEqual(report.errors, [])

    def test_placeholder_warnings_identify_each_source(self) -> None:
        normalized = normalize_fixture()
        warnings = validate(normalized).warnings
        self.assertEqual(len(warnings), 2)
        self.assertTrue(any("T-3" in warning and "task_role_effort" in warning for warning in warnings))
        self.assertTrue(any("T-4" in warning and "Task_dependencies" in warning for warning in warnings))

    def test_melt_effort_keeps_explicit_zero_and_missing_values(self) -> None:
        table = RawTable(
            name="effort",
            marker_row=None,
            header_row=1,
            headers=["Роль", "T-1", "T-2"],
            rows=[["Аналитик", 8, 0], ["Тестировщик", None, 12]],
        )
        rows, task_ids, roles = melt_effort(table)
        self.assertEqual(len(rows), 4)
        self.assertIn({"task_id": "T-1", "role": "Аналитик", "required_hours": 8}, rows)
        self.assertIn({"task_id": "T-2", "role": "Аналитик", "required_hours": 0}, rows)
        self.assertIn({"task_id": "T-1", "role": "Тестировщик", "required_hours": None}, rows)
        self.assertEqual(task_ids, {"T-1", "T-2"})
        self.assertEqual(roles, {"Аналитик", "Тестировщик"})

    def test_split_skills_preserves_parenthesized_commas(self) -> None:
        self.assertEqual(
            split_skills("Python, CI/CD (GitLab, Jenkins); PostgreSQL"),
            ["Python", "CI/CD (GitLab, Jenkins)", "PostgreSQL"],
        )


class ValidationTests(unittest.TestCase):
    def test_dependency_validation_reports_unknown_task_reference(self) -> None:
        issues = dependency_reference_issues(
            [{"blocking_task_id": "T-1", "blocked_task_id": "T-404"}], {"T-1"}
        )
        self.assertEqual(len(issues), 1)
        self.assertIn("T-404", issues[0])

    def test_capacity_rate_validation(self) -> None:
        self.assertEqual(validate_capacity_rate([{"engineer_id": "E-1", "capacity_rate": 0.5}]), [])
        violations = validate_capacity_rate([{"engineer_id": "E-2", "capacity_rate": 1.1}])
        self.assertEqual(len(violations), 1)


class IdempotencyTests(unittest.TestCase):
    def test_reference_lists_are_written_to_separate_csv_files(self) -> None:
        with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as temporary_directory:
            root = Path(temporary_directory)
            raw_dir = root / "data" / "raw"
            raw_dir.mkdir(parents=True)
            raw_file = raw_dir / "dataset.xlsx"
            build_fixture_workbook(raw_file)
            output_dir = root / "data" / "processed"
            run_preprocessing(raw_file, output_dir, root)
            with (output_dir / "planning_results.csv").open(
                encoding="utf-8", newline=""
            ) as source:
                planning_results = list(csv.DictReader(source))

        self.assertEqual(len(planning_results), 8)
        self.assertEqual(
            planning_results[-1],
            {"source_order": "8", "value": "Отменено зака"},
        )

    def test_repeat_run_skips(self) -> None:
        with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as temporary_directory:
            root = Path(temporary_directory)
            raw_dir = root / "data" / "raw"
            raw_dir.mkdir(parents=True)
            raw_file = raw_dir / "dataset.xlsx"
            build_fixture_workbook(raw_file)
            output_dir = root / "data" / "processed"
            first = run_preprocessing(raw_file, output_dir, root)
            second = run_preprocessing(raw_file, output_dir, root)
        self.assertEqual(first.status, "processed")
        self.assertEqual(second.status, "skipped")

    def test_force_rebuilds_processed_data(self) -> None:
        with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as temporary_directory:
            root = Path(temporary_directory)
            raw_dir = root / "data" / "raw"
            raw_dir.mkdir(parents=True)
            raw_file = raw_dir / "dataset.xlsx"
            build_fixture_workbook(raw_file)
            output_dir = root / "data" / "processed"
            first = run_preprocessing(raw_file, output_dir, root)
            forced = run_preprocessing(raw_file, output_dir, root, force=True)
        self.assertEqual(first.status, "processed")
        self.assertEqual(forced.status, "processed")


if __name__ == "__main__":
    unittest.main()
